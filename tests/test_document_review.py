import sys
import unittest
from io import BytesIO
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from agentops_api.document_review import DocumentInputError, extract_text, review_uploaded_deterministic  # noqa: E402


class DocumentReviewTests(unittest.TestCase):
    def test_text_upload_is_segmented_and_topic_tagged(self) -> None:
        text = (
            "The supplier's liability is capped at fees for six months.\n\n"
            "The supplier will process customer personal data to provide hosted services.\n\n"
            "The supplier will notify us promptly after a security incident."
        )
        extracted, clauses = extract_text("agreement.txt", text.encode())
        self.assertEqual(extracted, text)
        self.assertEqual(len(clauses), 3)
        self.assertEqual(clauses[0]["topic"], "Liability cap")

    def test_docx_extraction_includes_table_controls(self) -> None:
        from docx import Document

        document = Document()
        document.add_paragraph("Internal procurement requirements")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Supplier liability"
        table.cell(0, 1).text = "The liability cap must cover at least 12 months of fees."
        table.add_row().cells[0].text = "Security incidents"
        table.rows[1].cells[1].text = "Notify the company within 24 hours."
        content = BytesIO()
        document.save(content)

        extracted, clauses = extract_text("policy.docx", content.getvalue())

        self.assertIn("Supplier liability | The liability cap must cover at least 12 months of fees.", extracted)
        self.assertIn("Notify the company within 24 hours.", extracted)
        self.assertTrue(any("12 months" in clause["text"] for clause in clauses))

    def test_upload_rejects_unsupported_suffix_and_oversize(self) -> None:
        with self.assertRaisesRegex(DocumentInputError, "Unsupported file type"):
            extract_text("agreement.rtf", b"some text")
        with self.assertRaisesRegex(DocumentInputError, "10 MB"):
            extract_text("agreement.txt", b"x" * (10 * 1024 * 1024 + 1))

    def test_rules_upload_displays_all_three_evidence_classes(self) -> None:
        _, clauses = extract_text("agreement.txt", (
            "The supplier's liability is capped at fees for six months.\n\n"
            "The supplier will process customer personal data to provide hosted services.\n\n"
            "The supplier will notify us promptly after a security incident."
        ).encode())
        result = review_uploaded_deterministic("Demo Vendor", clauses, "excerpt")
        self.assertEqual(result["document_completeness"], "excerpt")
        self.assertEqual({item["evidence_status"] for item in result["findings"]}, {
            "confirmed_conflict", "missing_from_supplied_text", "requirement_not_established"
        })


if __name__ == "__main__":
    unittest.main()
